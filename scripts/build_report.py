"""Build the final DOCX report using generated metrics, plots, and evidence screenshots."""

from __future__ import annotations

import json
import os
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"
ARTIFACTS_DIR = PROJECT_ROOT / "artifacts"
SCREENSHOTS_DIR = PROJECT_ROOT / "screenshots"
OUTPUT_PATH = DOCS_DIR / "Heart_Disease_MLOps_Report.docx"


def read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size, color in (
        ("Title", 26, RGBColor(30, 70, 82)),
        ("Heading 1", 18, RGBColor(30, 70, 82)),
        ("Heading 2", 13, RGBColor(177, 75, 55)),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color


def add_title(document: Document) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("End-to-End MLOps for Heart Disease Risk Classification")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Assignment 01 | AIMLCZG523 Machine Learning Operations").bold = True
    document.add_paragraph("\n")
    details = document.add_table(rows=5, cols=2)
    details.style = "Light Shading Accent 1"
    values = (
        ("Student", os.getenv("STUDENT_NAME", "<ENTER STUDENT NAME>")),
        ("Student ID", os.getenv("STUDENT_ID", "<ENTER STUDENT ID>")),
        ("Dataset", "UCI Cleveland Heart Disease"),
        ("Repository", os.getenv("REPOSITORY_URL", "<ENTER CODE REPOSITORY URL>")),
        ("Report date", date.today().isoformat()),
    )
    for row, (label, value) in zip(details.rows, values, strict=True):
        row.cells[0].text = label
        row.cells[1].text = value
    document.add_paragraph("\n")
    statement = document.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.add_run(
        "This report documents a reproducible learning project. The API is an educational "
        "risk classifier and must not be used as a clinical diagnosis system."
    ).italic = True


def start_section(document: Document, title: str, number: int) -> None:
    document.add_page_break()
    document.add_heading(f"{number}. {title}", level=1)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_picture(document: Document, path: Path, caption: str, width: float = 6.4) -> None:
    if not path.is_file():
        document.add_paragraph(f"Evidence pending: generate {path.relative_to(PROJECT_ROOT)}")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True


def add_metric_table(document: Document, metadata: dict[str, Any]) -> None:
    candidates = metadata.get("candidates", [])
    if not candidates:
        document.add_paragraph(
            "Run the training command to populate this table with measured CV and test metrics."
        )
        return
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    headers = ("Model", "Split", "Accuracy", "Precision", "Recall", "ROC-AUC")
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for candidate in candidates:
        for split_name, metric_key in (("CV mean", "cv_metrics"), ("Test", "test_metrics")):
            metrics = candidate[metric_key]
            cells = table.add_row().cells
            values = (
                candidate["model_name"].replace("_", " ").title(),
                split_name,
                f"{metrics['accuracy']:.3f}",
                f"{metrics['precision']:.3f}",
                f"{metrics['recall']:.3f}",
                f"{metrics['roc_auc']:.3f}",
            )
            for cell, value in zip(cells, values, strict=True):
                cell.text = value


def build_report(output_path: Path = OUTPUT_PATH) -> Path:
    metadata = read_json(PROJECT_ROOT / "models" / "model_metadata.json")
    data_summary = read_json(PROJECT_ROOT / "data" / "processed" / "data_summary.json")
    document = Document()
    configure_document(document)
    document.core_properties.title = "Heart Disease MLOps Assignment Report"
    document.core_properties.subject = "End-to-end model development and deployment"
    add_title(document)

    start_section(document, "Executive Summary and Objectives", 1)
    document.add_paragraph(
        "This project implements the complete lifecycle of a binary heart disease risk model: "
        "versioned acquisition, data-quality checks, exploratory analysis, leakage-safe feature "
        "engineering, comparative model tuning, experiment tracking, package creation, automated "
        "testing, container delivery, Kubernetes deployment, and operational monitoring."
    )
    add_bullets(
        document,
        [
            (
                "Business objective: estimate the probability of heart disease from "
                "13 clinical attributes."
            ),
            (
                "Engineering objective: make every material artifact reproducible from "
                "a clean checkout."
            ),
            (
                "Selection objective: compare Logistic Regression and Random Forest by "
                "mean CV ROC-AUC."
            ),
            "Serving objective: return a validated prediction and confidence through FastAPI.",
            (
                "Operations objective: expose health, logs, Prometheus metrics, and a "
                "Grafana dashboard."
            ),
        ],
    )
    document.add_heading("Scope and responsible use", level=2)
    document.add_paragraph(
        "The solution is designed for coursework and technical demonstration. It does not "
        "establish clinical validity, calibrated treatment thresholds, fairness across protected "
        "groups, or regulatory compliance. A production healthcare use would require all of "
        "these controls."
    )

    start_section(document, "Data Acquisition and Quality", 2)
    document.add_paragraph(
        "The acquisition script downloads the processed Cleveland cohort from the UCI Machine "
        "Learning Repository. It retries transient failures, writes atomically, maps '?' tokens to "
        "missing values, validates domains, binarizes the original 0-4 diagnosis field, removes "
        "exact duplicates, and records a SHA-256 checksum and source URL."
    )
    data_table = document.add_table(rows=1, cols=2)
    data_table.style = "Light Shading Accent 1"
    data_table.rows[0].cells[0].text = "Quality measure"
    data_table.rows[0].cells[1].text = "Observed value"
    for label, value in (
        ("Rows", data_summary.get("rows", "Pending pipeline run")),
        ("Predictors", data_summary.get("feature_count", 13)),
        ("Duplicate rows after cleaning", data_summary.get("duplicate_rows", "Pending")),
        ("Class counts", str(data_summary.get("class_balance", "Pending"))),
        ("Raw SHA-256", data_summary.get("sha256", "Pending pipeline run")),
    ):
        cells = data_table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
    document.add_paragraph(
        "Missing predictors are intentionally not imputed in the persisted CSV. Imputers are "
        "fitted inside each cross-validation training fold, preventing information leakage."
    )

    start_section(document, "Exploratory Data Analysis", 3)
    document.add_paragraph(
        "The EDA job produces deterministic, high-resolution plots for class balance, "
        "missingness, numeric distributions by outcome, and the full Pearson correlation matrix. "
        "These views test whether imbalance, implausible values, or strongly redundant variables "
        "need intervention."
    )
    add_picture(
        document,
        ARTIFACTS_DIR / "eda" / "class_balance.png",
        "Figure 1. Binary target balance after cleaning.",
        width=5.7,
    )
    add_picture(
        document,
        ARTIFACTS_DIR / "eda" / "correlation_heatmap.png",
        "Figure 2. Correlation matrix for features and target.",
        width=5.7,
    )

    start_section(document, "Feature Engineering and Validation Design", 4)
    document.add_paragraph(
        "Five continuous variables use median imputation followed by standard scaling. Eight coded "
        "categorical variables use most-frequent imputation and one-hot encoding with "
        "unknown-category tolerance. A scikit-learn ColumnTransformer and classifier are saved "
        "as one pipeline, ensuring that training and inference execute identical transformations."
    )
    add_bullets(
        document,
        [
            "Train/test split: stratified 80/20 with random state 42.",
            "Tuning: shuffled stratified five-fold cross-validation on the training split only.",
            "Primary selection metric: mean validation ROC-AUC.",
            "Secondary metrics: accuracy, precision, recall, standard deviation, and fit time.",
            "Final test set: evaluated once for an unbiased estimate after candidate tuning.",
        ],
    )
    document.add_paragraph(
        "The linear model provides a lower-variance, interpretable baseline. The forest tests "
        "whether non-linear interactions and threshold effects improve ranking performance. "
        "Balanced class weights reduce sensitivity to moderate outcome imbalance."
    )

    start_section(document, "Experiments, Tuning, and Results", 5)
    document.add_paragraph(
        "Every candidate has a separate MLflow run containing the grid choice, cross-validation "
        "metrics and variability, held-out metrics, fit duration, ROC curve, confusion matrix, "
        "input example, model signature, and serialized pipeline."
    )
    add_metric_table(document, metadata)
    selected_model = metadata.get("selected_model", "Pending pipeline run")
    document.add_paragraph(
        f"Selected model: {selected_model}. Selection is based exclusively on mean cross-validated "
        "ROC-AUC, not on the test score. This rule is stored in model_metadata.json."
    )
    add_picture(
        document,
        ARTIFACTS_DIR / "training" / "model_comparison.png",
        "Figure 3. Cross-validation comparison used for model selection.",
        width=6.2,
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "03_mlflow_runs.png",
        "Figure 4. MLflow experiment comparison for both candidate models.",
        width=6.2,
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "04_mlflow_best_run.png",
        "Figure 5. Selected Logistic Regression run details in MLflow.",
        width=6.2,
    )

    start_section(document, "Packaging and Inference API", 6)
    document.add_paragraph(
        "The winning preprocessing-classifier pipeline is stored with Joblib and accompanied by "
        "JSON metadata that contains feature order, package version, candidate results, run "
        "identifier, "
        "and selection rationale. FastAPI validates all 13 fields and rejects missing, extra, or "
        "out-of-range values before invoking the model."
    )
    api_table = document.add_table(rows=1, cols=3)
    api_table.style = "Light Shading Accent 1"
    headers = ("Endpoint", "Purpose", "Evidence")
    for cell, value in zip(api_table.rows[0].cells, headers, strict=True):
        cell.text = value
    for values in (
        ("GET /health", "Readiness and model loading", "200 only when model is usable"),
        ("POST /predict", "Prediction and confidence", "Strict JSON schema"),
        ("GET /metrics", "Prometheus exposition", "Counts and latency histogram"),
        ("GET /docs", "OpenAPI user interface", "Interactive contract"),
    ):
        cells = api_table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    document.add_paragraph(
        "Structured logs include request ID, route, status, latency, and output class. Raw patient "
        "features are deliberately excluded from logs to reduce exposure of sensitive information."
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "07_docker_api.png",
        "Figure 6. Healthy isolated Docker container and successful prediction response.",
        width=6.2,
    )

    start_section(document, "Automated Testing and CI/CD", 7)
    document.add_paragraph(
        "Tests use deterministic synthetic data and cover schema failures, missing-value "
        "treatment, both model families, EDA generation, end-to-end MLflow packaging, API "
        "validation, model unavailability, and Prometheus output. The GitHub Actions workflow "
        "fails immediately on lint or test errors before running later stages."
    )
    add_bullets(
        document,
        [
            "Ruff enforces errors, imports, bugbear checks, and Python modernization rules.",
            "Pytest publishes JUnit and XML coverage evidence.",
            "Training uses a reduced but representative grid in CI to control execution time.",
            "The final model, MLflow store, plots, cleaned data, report, and logs are uploaded.",
            "Docker build proves that the serving environment is independently reproducible.",
        ],
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "06_ci_pipeline_success.png",
        "Figure 7. Successful CI/CD pipeline from lint through image packaging.",
        width=6.2,
    )

    start_section(document, "Containerization and Kubernetes Deployment", 8)
    document.add_paragraph(
        "The image uses a pinned Python 3.11 slim base, a serving-only dependency set, a non-root "
        "user, a model-aware health check, unbuffered logs, and one worker per container. Docker "
        "Compose adds read-only model mounting and a read-only API filesystem."
    )
    document.add_paragraph(
        "Kubernetes deploys two API replicas with readiness/liveness probes, CPU and memory "
        "requests and limits, dropped Linux capabilities, a read-only root filesystem, and a "
        "LoadBalancer service. Prometheus and Grafana share the namespace through Kustomize."
    )
    architecture = document.add_table(rows=3, cols=5)
    architecture.style = "Light Shading Accent 1"
    architecture.cell(0, 0).text = "UCI"
    architecture.cell(0, 1).text = "->"
    architecture.cell(0, 2).text = "Data/EDA"
    architecture.cell(0, 3).text = "->"
    architecture.cell(0, 4).text = "MLflow + Model"
    architecture.cell(1, 0).text = "GitHub"
    architecture.cell(1, 1).text = "->"
    architecture.cell(1, 2).text = "Actions"
    architecture.cell(1, 3).text = "->"
    architecture.cell(1, 4).text = "Docker Image"
    architecture.cell(2, 0).text = "Client"
    architecture.cell(2, 1).text = "->"
    architecture.cell(2, 2).text = "FastAPI/K8s"
    architecture.cell(2, 3).text = "->"
    architecture.cell(2, 4).text = "Prometheus/Grafana"
    add_picture(
        document,
        SCREENSHOTS_DIR / "09_kubernetes_workloads.png",
        "Figure 8. Healthy Kubernetes deployments, pods, and exposed services.",
        width=6.2,
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "10_kubernetes_predict.png",
        "Figure 9. Prediction served through the Kubernetes LoadBalancer service.",
        width=6.2,
    )

    start_section(document, "Monitoring, Logging, and Operations", 9)
    document.add_paragraph(
        "Prometheus scrapes the service every five seconds. The provisioned Grafana dashboard "
        "reports request rate by route, p95 latency, HTTP status trends, and output distribution. "
        "Health probes prevent traffic reaching a pod that cannot load its model. Request IDs "
        "support correlation from a client response to a JSON log line."
    )
    add_bullets(
        document,
        [
            "Availability signal: health probe and request status series.",
            "Latency signal: histogram-derived p95 by endpoint.",
            "Traffic signal: per-route request rate.",
            "Behavior signal: prediction class distribution for drift triage.",
            (
                "Recommended extension: validated feature drift, calibration drift, and "
                "outcome feedback."
            ),
        ],
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "05_prometheus_targets.png",
        "Figure 10. Prometheus API target in the UP state.",
        width=6.2,
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "11_grafana_dashboard.png",
        "Figure 11. Provisioned Grafana operations dashboard with live metrics.",
        width=6.2,
    )
    add_picture(
        document,
        SCREENSHOTS_DIR / "12_api_logs.png",
        "Figure 12. Structured request logs with status, latency, and request IDs.",
        width=6.2,
    )

    start_section(document, "Reproducibility, Limitations, and Conclusion", 10)
    document.add_paragraph(
        "Reproduction starts from pinned dependencies and a single documented command sequence. "
        "Source checksum, fixed random states, fold strategy, feature order, best hyperparameters, "
        "package version, MLflow run ID, and image tag make major transitions auditable."
    )
    document.add_heading("Limitations and next steps", level=2)
    add_bullets(
        document,
        [
            "The Cleveland cohort is small and may not represent current or global populations.",
            "No causal or clinical claim follows from predictive association.",
            "Confidence is model probability, not a guarantee of calibrated individual risk.",
            (
                "External validation, fairness analysis, threshold governance, and clinician "
                "review are required."
            ),
            (
                "Production should use a registry, signed images, secrets management, TLS, "
                "and durable metrics."
            ),
        ],
    )
    document.add_heading("Conclusion", level=2)
    document.add_paragraph(
        "The assignment demonstrates an integrated MLOps system rather than an isolated notebook. "
        "The same preprocessing artifact moves from cross-validation to API inference; CI "
        "recreates "
        "data, evidence, model, report, and image; Kubernetes supplies scalable delivery; and "
        "metrics close the operational feedback loop."
    )
    document.add_heading("Reproduction commands", level=2)
    for command in (
        "python -m heart_disease_mlops.data",
        "python -m heart_disease_mlops.eda",
        "python -m heart_disease_mlops.train",
        "pytest -q && ruff check src tests scripts",
        "docker compose up --build -d",
        "kubectl apply -k k8s",
    ):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    final_section = document.add_section(WD_SECTION.NEW_PAGE)
    add_page_number(final_section.footer.paragraphs[0])
    document.add_heading("Appendix: Submission Evidence", level=1)
    document.add_paragraph(
        "The repository screenshots directory contains execution evidence for MLflow, CI/CD, "
        "Docker, Kubernetes, Prometheus, Grafana, and structured API logs. The repository URL, "
        "local API access instructions, and demonstration video link complete the submission."
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    generated = build_report()
    print(f"Generated report: {generated}")
